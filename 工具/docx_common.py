# -*- coding: utf-8 -*-
"""共用 docx 工具：中文字型、表格、標題樣式"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "微軟正黑體"
INK = RGBColor(0x1F, 0x2E, 0x28)
ACC = RGBColor(0xB4, 0x5D, 0x22)
GREY = RGBColor(0x60, 0x6B, 0x65)


def set_run(run, size=12, bold=False, color=None, name=FONT):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(attr), name)


def new_doc(margin_cm=2.2):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(margin_cm)
    sec.bottom_margin = Cm(margin_cm)
    sec.left_margin = Cm(margin_cm)
    sec.right_margin = Cm(margin_cm)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.25
    return doc


def para(doc, text="", size=11, bold=False, color=None, align=None,
         space_before=0, space_after=4, style=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if text:
        r = p.add_run(text)
        set_run(r, size=size, bold=bold, color=color)
    return p


def h1(doc, text):
    para(doc, text, size=16, bold=True, color=INK, space_before=14, space_after=8)


def h2(doc, text):
    para(doc, text, size=12, bold=True, color=ACC, space_before=10, space_after=4)


def bullets(doc, items, size=11, numbered=False):
    for i, t in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.first_line_indent = Cm(-0.6)
        p.paragraph_format.space_after = Pt(3)
        prefix = f"{i}. " if numbered else "・"
        r = p.add_run(prefix)
        set_run(r, size=size)
        r2 = p.add_run(t)
        set_run(r2, size=size)


def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def table(doc, header, rows, widths=None, size=10.5, header_fill="E4EFE7"):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    tblPr = t._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    hdr = t.rows[0].cells
    # 表頭跨頁重複
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trPr.append(th)
    for i, h in enumerate(header):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(str(h))
        set_run(r, size=size, bold=True)
        shade(hdr[i], header_fill)
    for row in rows:
        r_obj = t.add_row()
        trPr2 = r_obj._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit")
        trPr2.append(cs)
        cells = r_obj.cells
        for i, v in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            r = p.add_run(str(v))
            set_run(r, size=size)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Cm(w)
    return t


def hr(doc):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:color"), "D8D2C4")
    pbdr.append(bottom)
    pPr.append(pbdr)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)
    return p


def save(doc, path):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    doc.save(path)
    print("已產生：", path)
