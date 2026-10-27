# -*- coding: utf-8 -*-
"""產生教案 docx（13 節，含 10 題隨堂測驗）"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from content import META, PLAN, QUIZ10
from docx_common import (new_doc, para, h1, h2, bullets, table, hr, save,
                         set_run, INK, ACC, GREY)

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "教案", "教案_視力口腔檢查站.docx")


def build():
    doc = new_doc()

    # 標題
    para(doc, "教學活動設計", size=11, color=GREY, align="center", space_after=2)
    para(doc, "視力口腔檢查站", size=26, bold=True, color=INK, align="center", space_after=2)
    para(doc, f"{META['publisher']}｜{META['unit']}", size=10.5, color=GREY,
         align="center", space_after=10)
    hr(doc)

    # 一、基本資料
    h1(doc, "一、基本資料")
    table(doc, ["項目", "內容"],
          [[k, v] for k, v in PLAN["basic"].items()],
          widths=[3.4, 13.0], size=10.5)

    # 二、設計理念
    h1(doc, "二、設計理念")
    bullets(doc, PLAN["idea"])

    # 三、學習重點
    h1(doc, "三、學習重點")
    table(doc, ["類別", "內容"],
          [[a, b] for a, b in PLAN["focus"]],
          widths=[3.4, 13.0], size=10.5)

    # 四、學習目標
    h1(doc, "四、學習目標")
    bullets(doc, PLAN["goals"], numbered=True)

    # 五、教材教具
    h1(doc, "五、教材教具")
    bullets(doc, PLAN["materials"])

    # 六、教學流程
    h1(doc, "六、教學流程（3 節，共 120 分鐘）")
    rows = []
    for lesson, mins, act, goal, assess in PLAN["flow"]:
        rows.append([f"第 {lesson} 節", f"{mins} 分", act, goal, assess])
    table(doc, ["節次", "時間", "教學活動", "學習目標", "評量"],
          rows, widths=[1.7, 1.5, 8.2, 2.3, 2.7], size=9.5)

    # 七、隨堂測驗（與網頁 10 題完全一致）
    h1(doc, "七、隨堂測驗（10 題，答案附於題後）")
    for i, q in enumerate(QUIZ10, 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt_(6)
        p.paragraph_format.space_after = Pt_(2)
        r = p.add_run(f"{i}. {q['q']}")
        set_run(r, size=11, bold=True)
        for k, opt in enumerate(q["o"]):
            op = doc.add_paragraph()
            op.paragraph_format.left_indent = Cm_(0.8)
            op.paragraph_format.space_after = Pt_(1)
            mark = "（✓）" if k == q["a"] else "（　）"
            rr = op.add_run(f"{mark}{'ABCD'[k]}. {opt}")
            set_run(rr, size=10.5, bold=(k == q["a"]))
        ep = doc.add_paragraph()
        ep.paragraph_format.left_indent = Cm_(0.8)
        ep.paragraph_format.space_after = Pt_(4)
        rr = ep.add_run(f"解析：{q['e']}")
        set_run(rr, size=10, color=GREY)

    # 八、常見迷思
    h1(doc, "八、常見迷思與澄清")
    table(doc, ["學生常這樣想", "澄清"],
          [[a, b] for a, b in PLAN["myths"]],
          widths=[6.2, 10.2], size=10.5)

    # 九、評量設計
    h1(doc, "九、評量設計")
    table(doc, ["類型", "方式", "比重"],
          [[a, b, c] for a, b, c in PLAN["assess"]],
          widths=[3.4, 10.6, 2.4], size=10.5)

    # 十、差異化教學
    h1(doc, "十、差異化教學")
    table(doc, ["學生需求", "調整方式"],
          [[a, b] for a, b in PLAN["diff"]],
          widths=[3.4, 13.0], size=10.5)

    # 十一、課後延伸
    h1(doc, "十一、課後延伸")
    bullets(doc, PLAN["extend"])

    # 十二、板書規劃
    h1(doc, "十二、板書規劃")
    for line in PLAN["board"]:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm_(0.6)
        p.paragraph_format.space_after = Pt_(3)
        r = p.add_run(line)
        set_run(r, size=10.5)

    # 十三、教學省思
    h1(doc, "十三、教學省思（課後填寫）")
    bullets(doc, PLAN["reflect"])
    para(doc, "實際狀況與調整：", size=10.5, color=GREY, space_before=6)
    for _ in range(3):
        para(doc, "＿" * 42, size=10.5, color=GREY, space_after=6)

    save(doc, OUT)


# 避免和 python-docx 的常數撞名
from docx.shared import Pt as _Pt, Cm as _Cm


def Pt_(v):
    return _Pt(v)


def Cm_(v):
    return _Cm(v)


if __name__ == "__main__":
    build()
