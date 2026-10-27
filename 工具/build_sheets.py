# -*- coding: utf-8 -*-
"""產生學習單 2 份與測驗卷 docx"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from content import META, QUIZ10, SHEET1, SHEET2
from docx_common import new_doc, para, h1, h2, table, hr, save, set_run, INK, ACC, GREY
from docx.shared import Pt, Cm
from docx.enum.text import WD_BREAK

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def build_sheet(data, out_name):
    doc = new_doc(margin_cm=2.0)

    para(doc, META["publisher"] + "｜" + META["unit"], size=9.5, color=GREY, align="center", space_after=2)
    para(doc, data["title"], size=20, bold=True, color=INK, align="center", space_after=2)
    para(doc, data["sub"], size=10.5, color=GREY, align="center", space_after=6)
    hr(doc)

    for heading, body in data["sections"]:
        if heading == "---PB---":
            doc.add_page_break()
            continue
        h2(doc, heading)
        if isinstance(body, dict):
            for ln in body.get("lines", []):
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = 1.4
                r = p.add_run(ln.replace("|", "｜"))
                set_run(r, size=11)
            if body.get("table"):
                header, rows = body["table"]
                table(doc, header, rows, widths=body.get("widths"), size=10.5)
            for ln in body.get("after", []):
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(3)
                r = p.add_run(ln.replace("|", "｜"))
                set_run(r, size=11)
            continue
        for ln in body:
            txt = ln.replace("|", "｜")
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.4
            r = p.add_run(txt)
            set_run(r, size=11)

    hr(doc)
    p = doc.add_paragraph()
    r = p.add_run("家長簽名：＿＿＿＿＿＿＿＿　　　　導師簽名：＿＿＿＿＿＿＿＿")
    set_run(r, size=10.5, color=GREY)
    p2 = doc.add_paragraph()
    r2 = p2.add_run(data["footer"])
    set_run(r2, size=9.5, color=GREY)

    save(doc, os.path.join(ROOT, "教案", out_name))


def build_quiz():
    doc = new_doc(margin_cm=2.0)

    para(doc, META["publisher"] + "｜" + META["unit"], size=9.5, color=GREY, align="center", space_after=2)
    para(doc, "隨堂測驗卷　視力口腔檢查站", size=20, bold=True, color=INK, align="center", space_after=2)
    para(doc, "五年級　座號：______　姓名：__________　　得分：__________",
         size=10.5, color=GREY, align="center", space_after=4)
    para(doc, "作答說明：每題只有一個正確答案，請把答案寫在題號前的括號內。共 10 題，每題 10 分。",
         size=10, color=GREY, space_after=6)
    hr(doc)

    for i, q in enumerate(QUIZ10, 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(7)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(f"（　　）{i}. {q['q']}")
        set_run(r, size=11.5, bold=True)
        for k, opt in enumerate(q["o"]):
            op = doc.add_paragraph()
            op.paragraph_format.left_indent = Cm(0.9)
            op.paragraph_format.space_after = Pt(1)
            rr = op.add_run(f"{'ABCD'[k]}. {opt}")
            set_run(rr, size=10.5)

    # 答案頁
    doc.add_page_break()
    para(doc, "答案與解析（教師用）", size=18, bold=True, color=INK, align="center", space_after=4)
    para(doc, "視力口腔檢查站｜10 題", size=10.5, color=GREY, align="center", space_after=6)
    hr(doc)

    ans = "　".join(f"{i+1}.{'ABCD'[q['a']]}" for i, q in enumerate(QUIZ10))
    p = doc.add_paragraph()
    r = p.add_run("答案一覽：" + ans)
    set_run(r, size=12, bold=True)
    p.paragraph_format.space_after = Pt(8)

    rows = []
    for i, q in enumerate(QUIZ10, 1):
        rows.append([f"{i}", "ABCD"[q["a"]], q.get("lv", ""), q["e"]])
    table(doc, ["題", "答案", "層次", "解析"], rows,
          widths=[1.0, 1.4, 1.6, 13.0], size=9.5)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    r = p.add_run("認知層次分布：記憶 3 題、理解 3 題、應用 3 題、分析 1 題。")
    set_run(r, size=10, color=GREY)
    p = doc.add_paragraph()
    r = p.add_run("若第 2、7、9 題錯誤率偏高，建議下節課用 5 分鐘複習「0.8 複檢線」與「3010」。")
    set_run(r, size=10, color=GREY)

    save(doc, os.path.join(ROOT, "教案", "測驗卷_視力口腔檢查站.docx"))


def main():
    build_sheet(SHEET1, "學習單1_護眼行動檢核表.docx")
    build_sheet(SHEET2, "學習單2_牙齒保衛戰.docx")
    build_quiz()


if __name__ == "__main__":
    main()
