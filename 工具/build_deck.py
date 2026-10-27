# -*- coding: utf-8 -*-
"""產生教學簡報 21 頁（16:9）→ 簡報/視力口腔檢查站_教學簡報.pptx"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from content import META, SLIDES

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "簡報", "視力口腔檢查站_教學簡報.pptx")

BG = RGBColor(0xFA, 0xF7, 0xF0)
INK = RGBColor(0x35, 0x54, 0x4A)
INK2 = RGBColor(0x6B, 0x7B, 0x72)
INK3 = RGBColor(0x9A, 0xA6, 0x9F)
ACC = RGBColor(0xC0, 0x6A, 0x2B)
SOFT = RGBColor(0xE4, 0xEF, 0xE7)
SOFT2 = RGBColor(0xF6, 0xE7, 0xD8)
LINE = RGBColor(0xE6, 0xDE, 0xCC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FONT = "微軟正黑體"
W, H = Inches(13.333), Inches(7.5)


def set_font(run, name=FONT):
    """python-pptx 只寫 a:latin，中文會掉字型，這裡補 a:ea。"""
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", name)


def textbox(slide, x, y, w, h, text, size=18, color=INK, bold=False,
            align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.25):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    lines = text.split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        set_font(r)
    return tb


def box(slide, x, y, w, h, fill=WHITE, line_col=LINE, radius=True, lw=1.0):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line_col is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line_col
        shp.line.width = Pt(lw)
    shp.shadow.inherit = False
    if radius:
        try:
            shp.adjustments[0] = 0.06
        except Exception:
            pass
    return shp


def new_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    bg.line.fill.background()
    bg.shadow.inherit = False
    return s


def header(slide, kicker, title, title_size=34):
    textbox(slide, Inches(0.9), Inches(0.55), Inches(11.5), Inches(0.35),
            kicker, size=14, color=ACC, bold=True)
    textbox(slide, Inches(0.9), Inches(0.95), Inches(11.5), Inches(0.9),
            title, size=title_size, color=INK, bold=True)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(1.92),
                                 Inches(1.05), Pt(4))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACC
    bar.line.fill.background()
    bar.shadow.inherit = False


def footer(slide, n):
    textbox(slide, Inches(0.9), Inches(6.92), Inches(8), Inches(0.3),
            f"{META['publisher']}｜{META['title']}", size=11, color=INK3)
    textbox(slide, Inches(12.0), Inches(6.92), Inches(0.6), Inches(0.3),
            f"{n:02d}", size=11, color=INK3, align=PP_ALIGN.RIGHT)


# ── 各版型 ──────────────────────────────────────────────────────

def cover(prs):
    d = SLIDES[0]
    s = new_slide(prs)
    band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.42), H)
    band.fill.solid()
    band.fill.fore_color.rgb = ACC
    band.line.fill.background()
    band.shadow.inherit = False

    textbox(s, Inches(1.1), Inches(1.35), Inches(11), Inches(0.4),
            d["kicker"], size=17, color=ACC, bold=True)
    textbox(s, Inches(1.1), Inches(1.95), Inches(11), Inches(1.6),
            d["title"], size=66, color=INK, bold=True)
    textbox(s, Inches(1.1), Inches(3.55), Inches(11), Inches(0.7),
            d["lead"], size=26, color=INK2)
    line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.15), Inches(4.45),
                              Inches(2.4), Pt(5))
    line.fill.solid()
    line.fill.fore_color.rgb = ACC
    line.line.fill.background()
    line.shadow.inherit = False

    x = Inches(1.1)
    for t in d["tags"]:
        w = Inches(0.32 * (len(t) + 2))
        box(s, x, Inches(4.95), w, Inches(0.52), fill=WHITE, line_col=LINE)
        textbox(s, x, Inches(5.06), w, Inches(0.3), t, size=15, color=INK2,
                align=PP_ALIGN.CENTER)
        x += w + Inches(0.16)

    textbox(s, Inches(1.1), Inches(6.1), Inches(11), Inches(0.4),
            "康軒版 健康與體育 五年級上學期｜3 節 120 分鐘", size=15, color=INK3)
    return s


def three(prs, d):
    s = new_slide(prs)
    header(s, d["kicker"], d["title"])
    n = len(d["bullets"])
    gap = Inches(0.32)
    total = Inches(11.53)
    cw = int((total - gap * (n - 1)) / n)
    x = Inches(0.9)
    for i, (head, body) in enumerate(d["bullets"]):
        box(s, x, Inches(2.35), cw, Inches(2.55), fill=WHITE, line_col=LINE)
        num = s.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(0.32), Inches(2.6),
                                 Inches(0.5), Inches(0.5))
        num.fill.solid()
        num.fill.fore_color.rgb = SOFT
        num.line.fill.background()
        num.shadow.inherit = False
        textbox(s, x + Inches(0.32), Inches(2.69), Inches(0.5), Inches(0.3),
                str(i + 1), size=15, color=ACC, bold=True, align=PP_ALIGN.CENTER)
        textbox(s, x + Inches(0.32), Inches(3.25), cw - Inches(0.64), Inches(0.6),
                head, size=21, color=INK, bold=True)
        textbox(s, x + Inches(0.32), Inches(3.88), cw - Inches(0.64), Inches(0.95),
                body, size=16, color=INK2, spacing=1.4)
        x += cw + gap
    if d.get("lead"):
        textbox(s, Inches(0.9), Inches(5.4), Inches(11.5), Inches(0.8),
                d["lead"], size=17, color=INK2, spacing=1.4)
    footer(s, d["n"])
    return s


def steps(prs, d):
    s = new_slide(prs)
    header(s, d["kicker"], d["title"])
    items = d["steps"]
    n = len(items)
    gap = Inches(0.26)
    total = Inches(11.53)
    cw = int((total - gap * (n - 1)) / n)
    x = Inches(0.9)
    for i, t in enumerate(items):
        box(s, x, Inches(2.7), cw, Inches(1.95), fill=WHITE if i % 2 else SOFT, line_col=None)
        textbox(s, x + Inches(0.28), Inches(2.92), cw - Inches(0.56), Inches(0.40),
                f"STEP {i + 1}", size=13, color=ACC, bold=True)
        textbox(s, x + Inches(0.28), Inches(3.36), cw - Inches(0.56), Inches(1.15),
                t, size=19, color=INK, bold=True, spacing=1.35)
        x += cw + gap
    if d.get("lead"):
        textbox(s, Inches(0.9), Inches(5.15), Inches(11.5), Inches(1.0),
                d["lead"], size=18, color=INK2, spacing=1.45)
    footer(s, d["n"])
    return s


def big(prs, d):
    s = new_slide(prs)
    header(s, d["kicker"], d["title"])
    box(s, Inches(0.9), Inches(2.45), Inches(11.53), Inches(2.35), fill=SOFT, line_col=None)
    textbox(s, Inches(0.9), Inches(2.72), Inches(11.53), Inches(1.2),
            d["big"], size=54, color=ACC, bold=True, align=PP_ALIGN.CENTER)
    textbox(s, Inches(1.55), Inches(3.95), Inches(10.2), Inches(0.7),
            d.get("lead", ""), size=18, color=INK2, align=PP_ALIGN.CENTER, spacing=1.4)
    footer(s, d["n"])
    return s


def main():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    for d in SLIDES:
        lay = d.get("layout")
        if lay == "cover":
            cover(prs)
        elif lay == "three":
            three(prs, d)
        elif lay == "steps":
            steps(prs, d)
        elif lay == "big":
            big(prs, d)
        else:
            raise ValueError("未知版型：" + str(lay))

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    prs.save(OUT)
    print(f"已產生 {len(prs.slides)} 頁簡報")
    print("輸出：", OUT)


if __name__ == "__main__":
    main()
